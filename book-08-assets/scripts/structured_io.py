"""Bounded CSV/JSON/XLSX readers plus atomic JSON output."""
from __future__ import annotations

import argparse
import csv
import json
import os
from pathlib import Path
from tempfile import NamedTemporaryFile
from typing import Any


class StructuredDataError(ValueError):
    pass


def _validate_rows(rows: list[dict[str, Any]], required: tuple[str, ...], max_rows: int) -> list[dict[str, Any]]:
    if len(rows) > max_rows:
        raise StructuredDataError("row_limit_exceeded")
    for index, row in enumerate(rows, start=1):
        if not isinstance(row, dict):
            raise StructuredDataError(f"row_{index}_not_object")
        missing = [name for name in required if name not in row]
        if missing:
            raise StructuredDataError(f"row_{index}_missing:{','.join(missing)}")
    return rows


def read_rows(path: str | Path, *, required: tuple[str, ...] = (), max_rows: int = 10_000) -> list[dict[str, Any]]:
    source = Path(path)
    if not source.is_file():
        raise StructuredDataError("file_not_found")
    suffix = source.suffix.casefold()
    if suffix == ".csv":
        with source.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            headers = reader.fieldnames or []
            if not headers or any(not name or not name.strip() for name in headers) or len(headers) != len(set(headers)):
                raise StructuredDataError("invalid_or_duplicate_csv_header")
            rows = []
            for row in reader:
                rows.append(dict(row))
                if len(rows) > max_rows:
                    raise StructuredDataError("row_limit_exceeded")
    elif suffix == ".json":
        try:
            payload = json.loads(source.read_text(encoding="utf-8-sig"))
        except (UnicodeDecodeError, json.JSONDecodeError) as exc:
            raise StructuredDataError("invalid_json") from exc
        rows = payload if isinstance(payload, list) else payload.get("records") if isinstance(payload, dict) else None
        if not isinstance(rows, list):
            raise StructuredDataError("json_must_be_array_or_records_object")
    elif suffix == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("install openpyxl to read XLSX") from exc
        workbook = load_workbook(source, read_only=True, data_only=True)
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        headers = [str(value).strip() if value is not None else "" for value in next(iterator, ())]
        if not headers or any(not name for name in headers) or len(headers) != len(set(headers)):
            raise StructuredDataError("invalid_or_duplicate_xlsx_header")
        rows = []
        for values in iterator:
            rows.append(dict(zip(headers, values, strict=False)))
            if len(rows) > max_rows:
                raise StructuredDataError("row_limit_exceeded")
        workbook.close()
    else:
        raise StructuredDataError("unsupported_structured_type")
    return _validate_rows(rows, required, max_rows)


def write_json_atomic(path: str | Path, payload: Any) -> None:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    temp_name = ""
    try:
        with NamedTemporaryFile("w", encoding="utf-8", newline="\n", dir=target.parent, delete=False, suffix=".tmp") as handle:
            json.dump(payload, handle, ensure_ascii=False, indent=2)
            handle.write("\n")
            handle.flush()
            os.fsync(handle.fileno())
            temp_name = handle.name
        os.replace(temp_name, target)
    finally:
        if temp_name:
            Path(temp_name).unlink(missing_ok=True)


def main() -> int:
    parser = argparse.ArgumentParser(description="Read bounded CSV, JSON or XLSX rows")
    parser.add_argument("path")
    parser.add_argument("--required", nargs="*", default=[])
    parser.add_argument("--max-rows", type=int, default=10_000)
    args = parser.parse_args()
    try:
        rows = read_rows(args.path, required=tuple(args.required), max_rows=args.max_rows)
    except (StructuredDataError, RuntimeError) as exc:
        print(json.dumps({"ok": False, "error": str(exc)}))
        return 2
    print(json.dumps({"ok": True, "rows": len(rows), "fields": sorted(rows[0]) if rows else []}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
